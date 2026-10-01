"""Synthetic end-to-end installation probe; never a real-data acceptance."""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from pypdf import PdfReader

from pseudokrat.ki_office import ProjectError
from pseudokrat.local_report import generate_report
from pseudokrat.report_preview import render_preview


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_smoke(output: Path) -> Path:
    """Create only synthetic inputs and check the native Calc-to-Writer path."""
    if platform.system() != "Linux":
        raise ProjectError("Der Spark-Selbsttest benötigt Linux, LibreOffice und UNO.")
    try:
        output = output.resolve()
        output.mkdir(mode=0o700, parents=True, exist_ok=False)
        excel, template, mapping = (
            output / name for name in ("synthetic.xlsx", "template.docx", "mapping.json")
        )
        book = Workbook()
        sheet = book.active
        if sheet is None:
            raise ProjectError("Synthetische Arbeitsmappe konnte nicht erstellt werden.")
        sheet.title = "Daten"
        sheet.append(["Name", "Saldo EUR"])
        sheet.append(["Testfall Alpha", "=ROUND(12.345*2,2)"])
        sheet.append(["Testfall Beta", "=-B2"])
        book.save(excel)
        doc = Document()
        footer = doc.sections[0].footer.paragraphs[0]
        footer.add_run("Seite ")
        for code, suffix in (("PAGE", " von "), ("NUMPAGES", "")):
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), code)
            run, text = OxmlElement("w:r"), OxmlElement("w:t")
            text.text = "999"
            run.append(text)
            field.append(run)
            footer._p.append(field)
            footer.add_run(suffix)
        for index in range(65):
            if index:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(f"Synthetischer Abschnitt {index + 1}", level=1)
            doc.add_paragraph("{{ assessment }}")
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text, row.cells[1].text = "{{ rows.name }}", "{{ rows.amount }} EUR"
        doc.save(str(template))
        spec = {
            "version": 3,
            "reviewed": True,
            "template_sha256": _hash(template),
            "headers": {"Daten": {"A1": "Name", "B1": "Saldo EUR"}},
            "fields": {
                "amount": {
                    "sheet": "Daten",
                    "range": "B2:B3",
                    "operation": "sum",
                    "format": "decimal",
                    "decimals": 2,
                }
            },
            "narratives": {
                "assessment": {"template": "Der synthetische Gesamtsaldo beträgt {{ amount }} EUR."}
            },
            "tables": {
                "rows": {
                    "sheet": "Daten",
                    "first_row": 2,
                    "last_row": 3,
                    "columns": {
                        "name": {"column": "A", "format": "text"},
                        "amount": {"column": "B", "format": "decimal", "decimals": 2},
                    },
                }
            },
        }
        mapping.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
        inputs = {path.name: _hash(path) for path in (excel, template, mapping)}
        report = generate_report(
            excel, template, mapping, output / "report", recalculate_libreoffice=True
        )
        rendered = render_preview(report / "bericht.docx", output / "preview")
        result = Document(str(report / "bericht.docx"))
        expected_text = "Der synthetische Gesamtsaldo beträgt 0,00 EUR."
        rows = [[cell.text for cell in row.cells] for row in result.tables[0].rows]
        reader = PdfReader(rendered / "preview.pdf")
        first = " ".join(reader.pages[0].extract_text().split())
        last = " ".join(reader.pages[-1].extract_text().split())
        checks = {
            "inputs_unchanged": inputs
            == {path.name: _hash(path) for path in (excel, template, mapping)},
            "narratives": sum(p.text == expected_text for p in result.paragraphs) == 65,
            "table": rows == [["Testfall Alpha", "24,69 EUR"], ["Testfall Beta", "-24,69 EUR"]],
            "pages": len(reader.pages) == 65,
            "page_fields": "Seite 1 von 65" in first and "Seite 65 von 65" in last,
        }
        evidence = {
            "synthetic_only": True,
            "production_approved": False,
            "status": "passed" if all(checks.values()) else "failed",
            "checks": checks,
            "system": platform.system(),
            "architecture": platform.machine(),
            "python": platform.python_version(),
            "input_sha256": inputs,
            "report_sha256": _hash(report / "bericht.docx"),
            "pdf_sha256": _hash(rendered / "preview.pdf"),
        }
        (output / "smoke.json").write_text(json.dumps(evidence, indent=2), encoding="utf-8")
        if not all(checks.values()):
            raise ProjectError("Synthetischer Selbsttest fehlgeschlagen; lokale Ergebnisse prüfen.")
        return output
    except ProjectError:
        raise
    except (OSError, ValueError, KeyError, IndexError, TypeError) as exc:
        raise ProjectError(
            "Spark-Selbsttest fehlgeschlagen; Installation und Ausgabe lokal prüfen."
        ) from exc


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Synthetischer 65-Seiten-Selbsttest mit lokalem Calc und Writer."
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        run_smoke(args.output)
        print("Synthetischer Selbsttest bestanden. Echte Dateien und Layout noch nicht abgenommen.")
        return 0
    except ProjectError:
        print(
            "Selbsttest fehlgeschlagen. Lokale Installation und neuen Ausgabeordner prüfen.",
            file=sys.stderr,
        )
        return 20


if __name__ == "__main__":
    raise SystemExit(main())

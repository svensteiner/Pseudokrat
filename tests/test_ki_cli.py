"""Public CLI journey without real documents or network access."""

from docx import Document
from openpyxl import Workbook

from pseudokrat.cli import main


def test_cli_prepare_publish_restore_and_local_handoff(tmp_path, data_dir, capsys):
    book = Workbook()
    book.active.append(["Name", "Betrag"])
    book.active.append(["Sondername", 500.25])
    source = tmp_path / "input.xlsx"
    book.save(source)
    word = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("Sondername")
    doc.save(word)
    terms = tmp_path / "terms.txt"
    terms.write_text("Sondername", encoding="utf-8")
    project = tmp_path / "project"
    common = ["--profile", "test", "--password", "synthetic-test-password"]
    assert main(["ki-project", *common, "prepare", "--input", str(source), str(word),
                 "--output", str(project), "--terms", str(terms)]) == 0
    assert main(["ki-project", *common, "publish", "--project", str(project)]) != 0
    assert main(["ki-project", *common, "publish", "--project", str(project),
                 "--reviewed", "--accept-numeric-linkability"]) == 0
    assert main(["ki-project", *common, "restore", "--project", str(project),
                 "--input", str(project / "vorschau" / "daten_01.xlsx"),
                 "--output", str(tmp_path / "restored.xlsx")]) == 0
    assert main(["ki-project", *common, "local-workspace", "--project", str(project),
                 "--output", str(tmp_path / "spark")]) == 0
    assert "Sondername" not in capsys.readouterr().out

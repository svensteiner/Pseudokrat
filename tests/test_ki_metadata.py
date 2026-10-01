"""Provenance must be removed without relying on name recognition."""

from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from lxml import etree
from openpyxl import Workbook, load_workbook

from pseudokrat.ki_office import ProjectError, S, read_office, write_office
from pseudokrat.ki_project import prepare_project, publish_project, restore_file


@pytest.mark.parametrize("suffix", ["xlsx", "docx"])
def test_export_removes_creator_and_hidden_provenance(tmp_path, store_and_audit, suffix):
    marker = "OpaqueOriginZXQ987"
    original = tmp_path / f"source.{suffix}"
    if suffix == "xlsx":
        book = Workbook()
        book.active["A1"] = 42.25
        book.active["B1"] = "=A1*2"
        book.properties.creator = marker
        book.properties.lastModifiedBy = marker
        book.save(original)
    else:
        doc = Document()
        doc.add_paragraph("Betrag 42,25 EUR")
        doc.core_properties.author = marker
        doc.core_properties.last_modified_by = marker
        doc.save(original)
    source = tmp_path / f"metadata.{suffix}"
    with ZipFile(original) as archive, ZipFile(source, "w", ZIP_DEFLATED) as target:
        target.comment = marker.encode()
        for name in archive.namelist():
            data = archive.read(name)
            if name.endswith(".xml") or name.endswith(".rels"):
                root = etree.fromstring(data)
                if name.startswith("docProps/"):
                    root.set("origin", marker)
                    root.text = marker
                if name == "xl/workbook.xml":
                    etree.SubElement(
                        root, f"{{{S}}}fileSharing", userName=marker, readOnlyRecommended="1"
                    )
                    etree.SubElement(root, f"{{{S}}}fileVersion", appName=marker, codeName=marker)
                    root.find(f"{{{S}}}workbookPr").set("codeName", marker)
                if name == "xl/worksheets/sheet1.xml":
                    root.find(f"{{{S}}}sheetPr").set("codeName", marker)
                if name == "[Content_Types].xml":
                    etree.SubElement(
                        root,
                        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                        PartName="/docProps/custom.xml",
                        ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml",
                    )
                if name == "_rels/.rels":
                    etree.SubElement(
                        root,
                        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                        Id="rIdCustom",
                        Target="docProps/custom.xml",
                        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties",
                    )
                data = etree.tostring(root)
            target.writestr(name, data)
        target.writestr(
            "docProps/custom.xml",
            f'<Properties origin="{marker}"><Company>{marker}</Company></Properties>',
        )
    before = source.read_bytes()
    key = store_and_audit[0].keys.fernet
    project = prepare_project([source], tmp_path / "project", key)
    published = publish_project(project, key, reviewed=True, accept_numeric_linkability=True)
    file = f"daten_01.{suffix}" if suffix == "xlsx" else "vorlage_01.docx"
    preview = project / "vorschau" / file
    with ZipFile(preview) as archive:
        assert not archive.comment
        assert not any(n.startswith("docProps/") for n in archive.namelist())
        assert marker.encode() not in b"".join(archive.read(n) for n in archive.namelist())
        assert b"docProps/" not in archive.read("[Content_Types].xml")
        assert b"docProps/" not in archive.read("_rels/.rels")
        if suffix == "xlsx":
            root = etree.fromstring(archive.read("xl/workbook.xml"))
            sharing = root.find(f"{{{S}}}fileSharing")
            assert sharing.get("userName") is None
            assert sharing.get("readOnlyRecommended") == "1"
    with ZipFile(published) as archive:
        assert archive.read(file) == preview.read_bytes()
    restored = restore_file(project, key, preview, tmp_path / f"restored.{suffix}")
    if suffix == "xlsx":
        result = load_workbook(restored)
        assert result.active["A1"].value == 42.25
        assert result.active["B1"].value == "=A1*2"
    else:
        assert Document(restored).paragraphs[0].text == "Betrag 42,25 EUR"
    assert source.read_bytes() == before


@pytest.mark.parametrize(
    "tag,part",
    [
        ("customWorkbookViews", "xl/workbook.xml"),
        ("customSheetViews", "xl/worksheets/sheet1.xml"),
    ],
)
def test_custom_views_with_provenance_are_blocked(tmp_path, tag, part):
    source = tmp_path / "views.xlsx"
    Workbook().save(source)
    roots = read_office(source)
    etree.SubElement(roots[part], f"{{{S}}}{tag}")
    source.write_bytes(write_office(roots))
    with pytest.raises(ProjectError, match=tag):
        read_office(source)

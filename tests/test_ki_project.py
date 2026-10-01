"""Real Office files, entirely artificial content, exercise the privacy boundary."""

import json
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from pseudokrat.ki_project import (
    ProjectError,
    prepare_original_workspace,
    prepare_project,
    publish_project,
    restore_file,
)


@pytest.fixture
def inputs(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Geheimfirma"
    sheet.append(["Kundennummer", "Name", "Betrag", "Menge"])
    sheet.append([123456, "Sondername", 1234.56, 3])
    sheet.append([123457, "Anderername", -25.5, 2])
    sheet["E2"] = "=C2*D2"
    second = book.create_sheet("Auswertung")
    second["A1"] = "='Geheimfirma'!E2"
    book.properties.creator = "Geheimfirma"
    xlsx = tmp_path / "Geheimfirma.xlsx"
    book.save(xlsx)
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Sonder").bold = True
    p.add_run("name erhält einen Bericht.")
    doc.add_paragraph("Betrag 1234,56 EUR")
    doc.core_properties.author = "Geheimfirma"
    word = tmp_path / "Geheimfirma.docx"
    doc.save(word)
    return xlsx, word


def prepare(tmp_path, inputs, store_and_audit):
    store, _ = store_and_audit
    return prepare_project(
        inputs,
        tmp_path / "project",
        store.keys.fernet,
        terms=["Geheimfirma", "Sondername", "Anderername"],
    )


def test_package_preserves_numbers_formulas_and_project_consistency(
    tmp_path, inputs, store_and_audit
):
    project = prepare(tmp_path, inputs, store_and_audit)
    assert not (project / "KI-Paket.zip").exists()
    book = load_workbook(project / "vorschau" / "daten_01.xlsx")
    assert book.worksheets[0]["C2"].value == 1234.56
    assert book.worksheets[0]["C3"].value == -25.5
    assert book.worksheets[0]["E2"].value == "=C2*D2"
    assert book.worksheets[1]["A1"].value == "='Blatt_001'!E2"
    assert book.worksheets[0]["A2"].value != 123456
    name = book.worksheets[0]["B2"].value
    doc = Document(project / "vorschau" / "vorlage_02.docx")
    assert name in doc.paragraphs[0].text
    assert doc.paragraphs[1].text == "Betrag 1234,56 EUR"
    assert b"Sondername" not in (project / "zuordnung.enc").read_bytes()


def test_large_word_template_roundtrip_with_tables(tmp_path, store_and_audit):
    """65 explicit page sections; this checks content, not Word pagination."""
    source = tmp_path / "large.docx"
    document = Document()
    for index in range(65):
        if index:
            document.add_page_break()
        document.add_heading(f"Abschnitt {index + 1}", level=1)
        paragraph = document.add_paragraph()
        paragraph.add_run("Sonder").bold = True
        paragraph.add_run("name: Auswertung")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Betrag"
        table.cell(1, 0).text = "Sondername"
        table.cell(1, 1).text = f"{index + 1},50 EUR"
    document.save(source)
    key = store_and_audit[0].keys.fernet
    project = prepare_project([source], tmp_path / "large-project", key, terms=["Sondername"])
    preview = project / "vorschau" / "vorlage_01.docx"
    masked = Document(preview)
    assert len(masked.tables) == 65
    assert all("Sondername" not in table.cell(1, 0).text for table in masked.tables)
    restored = tmp_path / "restored.docx"
    restore_file(project, key, preview, restored)
    result = Document(restored)
    for index, table in enumerate(result.tables):
        assert table.cell(1, 0).text == "Sondername"
        assert table.cell(1, 1).text == f"{index + 1},50 EUR"


def test_identity_hidden_in_number_format_blocks_export(tmp_path, store_and_audit):
    book = Workbook()
    book.active["A1"] = 123
    book.active["A1"].number_format = '0 "secret.person@example.com"'
    source = tmp_path / "format.xlsx"
    book.save(source)
    with pytest.raises(ProjectError):
        prepare_project([source], tmp_path / "blocked", store_and_audit[0].keys.fernet)
    assert not (tmp_path / "blocked").exists()


def test_release_requires_review_and_explicit_numeric_risk_acceptance(
    tmp_path, inputs, store_and_audit
):
    project = prepare(tmp_path, inputs, store_and_audit)
    key = store_and_audit[0].keys.fernet
    with pytest.raises(ProjectError, match="Prüfung"):
        publish_project(project, key)
    with pytest.raises(ProjectError, match="Zahlen"):
        publish_project(project, key, reviewed=True)
    output = publish_project(project, key, reviewed=True, accept_numeric_linkability=True)
    with ZipFile(output) as archive:
        assert "zuordnung.enc" not in archive.namelist()
        for name in archive.namelist():
            raw = archive.read(name)
            assert b"Geheimfirma" not in raw
        assert "entwicklungsauftrag.md" in archive.namelist()
    assert json.loads((project / "vorschau" / "schema.json").read_text("utf-8"))[
        "numbers_preserved"
    ]


def test_modified_preview_cannot_be_published(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    (project / "vorschau" / "schema.json").write_text("secret", encoding="utf-8")
    with pytest.raises(ProjectError, match="verändert"):
        publish_project(
            project, store_and_audit[0].keys.fernet, reviewed=True, accept_numeric_linkability=True
        )


def test_restore_resolves_split_word_and_numeric_identifiers(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    key = store_and_audit[0].keys.fernet
    result = tmp_path / "restored.xlsx"
    restore_file(project, key, project / "vorschau" / "daten_01.xlsx", result)
    book = load_workbook(result)
    assert book.worksheets[0]["A2"].value == 123456
    assert book.worksheets[0]["B2"].value == "Sondername"
    assert book.sheetnames[0] == "Geheimfirma"
    assert book.worksheets[1]["A1"].value == "='Geheimfirma'!E2"
    doc_result = tmp_path / "restored.docx"
    restore_file(project, key, project / "vorschau" / "vorlage_02.docx", doc_result)
    assert Document(doc_result).paragraphs[0].text == "Sondername erhält einen Bericht."


def test_unsupported_office_channels_fail_closed(tmp_path, inputs, store_and_audit):
    with ZipFile(inputs[0], "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/embeddings/secret.bin", b"secret")
    with pytest.raises(ProjectError, match="unterstützt"):
        prepare(tmp_path, inputs, store_and_audit)
    assert not (tmp_path / "project").exists()


def test_unknown_tokens_do_not_get_guessed(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    file = tmp_path / "answer.docx"
    doc = Document()
    doc.add_paragraph("[[PK_unknown_00001]]")
    doc.save(file)
    with pytest.raises(ProjectError, match="Platzhalter"):
        restore_file(project, store_and_audit[0].keys.fernet, file, tmp_path / "out.docx")
    assert not (tmp_path / "out.docx").exists()


def test_existing_destinations_are_never_overwritten(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    with pytest.raises(ProjectError):
        prepare_project(inputs, project, store_and_audit[0].keys.fernet)


def test_duplicate_basename_and_shared_identifiers(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    assert (project / "vorschau" / "daten_01.xlsx").is_file()
    assert (project / "vorschau" / "vorlage_02.docx").is_file()


def test_original_workspace_for_spark_uses_real_values_with_development_sheet_names(
    tmp_path, inputs, store_and_audit
):
    project = prepare(tmp_path, inputs, store_and_audit)
    result = prepare_original_workspace(
        project, store_and_audit[0].keys.fernet, tmp_path / "local-only"
    )
    book = load_workbook(result / "daten_01.xlsx")
    assert book.sheetnames[0] == "Blatt_001"
    assert book.worksheets[0]["B2"].value == "Sondername"
    assert book.worksheets[0]["C2"].value == 1234.56
    assert (result / "SPARK-AUFTRAG.md").exists()
    assert not (result / "KI-Paket.zip").exists()


def test_changed_originals_require_new_project(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    book = load_workbook(inputs[0])
    book.active["C2"] = 77
    book.save(inputs[0])
    with pytest.raises(ProjectError, match="Original"):
        prepare_original_workspace(project, store_and_audit[0].keys.fernet, tmp_path / "local-only")


def test_metadata_and_thumbnail_are_absent_from_preview(tmp_path, inputs, store_and_audit):
    project = prepare(tmp_path, inputs, store_and_audit)
    with ZipFile(project / "vorschau" / "vorlage_02.docx") as archive:
        assert not any(n.startswith("customXml/") or "thumbnail" in n for n in archive.namelist())
        assert b"Geheimfirma" not in b"".join(archive.read(n) for n in archive.namelist())


def test_external_link_blocks_export(tmp_path, inputs, store_and_audit):
    book = load_workbook(inputs[0])
    book.active["F1"].hyperlink = "https://example.org/private"
    book.save(inputs[0])
    with pytest.raises(ProjectError):
        prepare(tmp_path, inputs, store_and_audit)


def test_formula_criteria_replaced_consistently(tmp_path, inputs, store_and_audit):
    book = load_workbook(inputs[0])
    book.active["F2"] = '=SUMIF(B2:B3,"Sondername",C2:C3)'
    book.save(inputs[0])
    project = prepare(tmp_path, inputs, store_and_audit)
    clean = load_workbook(project / "vorschau" / "daten_01.xlsx")
    assert clean.active["F2"].value == f'=SUMIF(B2:B3,"{clean.active["B2"].value}",C2:C3)'


def test_hidden_sheet_is_processed(tmp_path, inputs, store_and_audit):
    book = load_workbook(inputs[0])
    sheet = book.create_sheet("Versteckt")
    sheet["A1"] = "Sondername"
    sheet.sheet_state = "veryHidden"
    book.save(inputs[0])
    project = prepare(tmp_path, inputs, store_and_audit)
    clean = load_workbook(project / "vorschau" / "daten_01.xlsx")
    assert clean.worksheets[2]["A1"].value == clean.active["B2"].value


def test_foreign_key_cannot_unlock_vault(tmp_path, inputs, store_and_audit):
    from cryptography.fernet import Fernet

    project = prepare(tmp_path, inputs, store_and_audit)
    with pytest.raises(ProjectError, match="Profil"):
        publish_project(
            project, Fernet(Fernet.generate_key()), reviewed=True, accept_numeric_linkability=True
        )


@pytest.mark.parametrize(
    "formula", ["=COUNT(A2:A3)", "=SUM(A:A)", "=A2+1", "=COUNTIF(A2:A3,123456)"]
)
def test_formulas_using_numeric_identifiers_block_instead_of_changing_results(
    tmp_path, inputs, store_and_audit, formula
):
    book = load_workbook(inputs[0])
    book.active["F2"] = formula
    book.save(inputs[0])
    with pytest.raises(ProjectError, match="Kennung"):
        prepare(tmp_path, inputs, store_and_audit)


def test_numeric_id_does_not_replace_unrelated_word_amount(tmp_path, inputs, store_and_audit):
    doc = Document(inputs[1])
    doc.add_paragraph("Betrag 123456 EUR")
    doc.save(inputs[1])
    project = prepare(tmp_path, inputs, store_and_audit)
    assert (
        Document(project / "vorschau" / "vorlage_02.docx").paragraphs[-1].text
        == "Betrag 123456 EUR"
    )


def test_validation_list_pii_is_discovered_without_manual_terms(tmp_path, store_and_audit):
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    validation = DataValidation(
        type="list", formula1='"secret.person@example.com,other@example.com"'
    )
    validation.add("A1")
    book.active.add_data_validation(validation)
    source = tmp_path / "validation.xlsx"
    book.save(source)
    project = prepare_project([source], tmp_path / "project", store_and_audit[0].keys.fernet)
    with ZipFile(project / "vorschau" / "daten_01.xlsx") as archive:
        assert b"secret.person@example.com" not in archive.read("xl/worksheets/sheet1.xml")


def test_equal_string_and_numeric_identifiers_keep_distinct_types(tmp_path, store_and_audit):
    book = Workbook()
    book.active.append(["Kundennummer", "Name"])
    book.active.append([123456, "123456"])
    source = tmp_path / "types.xlsx"
    book.save(source)
    key = store_and_audit[0].keys.fernet
    project = prepare_project([source], tmp_path / "project", key)
    restored = restore_file(
        project, key, project / "vorschau" / "daten_01.xlsx", tmp_path / "restored.xlsx"
    )
    result = load_workbook(restored).active
    assert result["A2"].value == 123456
    assert result["A2"].data_type == "n"
    assert result["B2"].value == "123456"
    assert result["B2"].data_type == "s"


def test_phone_numbers_still_replace_in_prose(tmp_path, store_and_audit):
    doc = Document()
    doc.add_paragraph("Telefon +43 664 1234567")
    source = tmp_path / "phone.docx"
    doc.save(source)
    project = prepare_project([source], tmp_path / "project", store_and_audit[0].keys.fernet)
    assert (
        "+43 664 1234567"
        not in Document(project / "vorschau" / "vorlage_01.docx").paragraphs[0].text
    )


def test_sensitive_cached_formula_result_blocks_reconstruction(tmp_path, store_and_audit):
    import io

    from lxml import etree

    book = Workbook()
    book.active["A1"] = '=CONCAT("secret",".person","@example.com")'
    buffer = io.BytesIO()
    book.save(buffer)
    source = tmp_path / "cache.xlsx"
    with ZipFile(buffer) as archive, ZipFile(source, "w", ZIP_DEFLATED) as target:
        for name in archive.namelist():
            data = archive.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = etree.fromstring(data)
                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                cell = next(root.iter(ns + "c"))
                cell.set("t", "str")
                cell.find(ns + "v").text = "secret.person@example.com"
                data = etree.tostring(root)
            target.writestr(name, data)
    with pytest.raises(ProjectError, match="Berechnet"):
        prepare_project([source], tmp_path / "project", store_and_audit[0].keys.fernet)

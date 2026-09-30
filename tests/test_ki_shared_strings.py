"""Excel string tables must not retain orphaned identifying content."""

from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook

from pseudokrat.ki_office import ProjectError, S, cell_text
from pseudokrat.ki_project import prepare_project, restore_file


def test_orphan_shared_strings_are_removed_after_identity_replacement(tmp_path, store_and_audit):
    book = Workbook()
    book.active.append(["Kundennummer", "Betrag"])
    book.active.append(["00000001", 1234.56])
    book.active.append([None, "Visible ordinary label"])
    original = tmp_path / "inline.xlsx"
    book.save(original)
    source = tmp_path / "shared.xlsx"
    with ZipFile(original) as archive, ZipFile(source, "w", ZIP_DEFLATED) as target:
        for name in archive.namelist():
            data = archive.read(name)
            root = etree.fromstring(data)
            if name == "xl/worksheets/sheet1.xml":
                cell = next(c for c in root.iter(f"{{{S}}}c") if c.get("r") == "A2")
                for child in list(cell):
                    cell.remove(child)
                cell.set("t", "s")
                etree.SubElement(cell, f"{{{S}}}v").text = "0"
                remaining = next(c for c in root.iter(f"{{{S}}}c") if c.get("r") == "B3")
                for child in list(remaining):
                    remaining.remove(child)
                remaining.set("t", "s")
                etree.SubElement(remaining, f"{{{S}}}v").text = "2"
            elif name == "[Content_Types].xml":
                etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                                 PartName="/xl/sharedStrings.xml",
                                 ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
            elif name == "xl/_rels/workbook.xml.rels":
                etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                                 Id="rIdStrings", Target="sharedStrings.xml",
                                 Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings")
            target.writestr(name, etree.tostring(root))
        target.writestr("xl/sharedStrings.xml", (
            f'<sst xmlns="{S}" count="2" uniqueCount="3"><si><t>00000001</t></si>'
            '<si><t>OrphanProjectOrchid</t></si><si><t>Visible ordinary label</t></si></sst>'
        ))
    key = store_and_audit[0].keys.fernet
    project = prepare_project([source], tmp_path / "project", key)
    preview = project / "vorschau" / "daten_01.xlsx"
    with ZipFile(preview) as archive:
        raw = archive.read("xl/sharedStrings.xml")
        assert b"00000001" not in raw
        assert b"OrphanProjectOrchid" not in raw
        table = etree.fromstring(raw)
        assert table.get("count") == "1"
        assert table.get("uniqueCount") == "1"
    assert load_workbook(preview).active["B3"].value == "Visible ordinary label"
    restored = restore_file(project, key, preview, tmp_path / "restored.xlsx")
    result = load_workbook(restored).active
    assert result["A2"].value == "00000001"
    assert result["A2"].data_type == "s"
    assert result["B2"].value == 1234.56
    assert result["B3"].value == "Visible ordinary label"


def test_negative_shared_string_index_is_invalid():
    cell = etree.fromstring(f'<c xmlns="{S}" t="s"><v>-1</v></c>')
    with pytest.raises(ProjectError, match="Textverweis"):
        cell_text(cell, ["last entry"])

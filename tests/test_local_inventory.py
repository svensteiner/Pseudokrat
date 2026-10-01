"""Local discovery is evidence for mapping, never a data-export approval."""

import json

from docx import Document
from openpyxl import Workbook

from pseudokrat.local_inventory import inspect_documents, main


def test_inventory_locates_sources_templates_and_missing_formula_cache(tmp_path):
    excel = tmp_path / "private.xlsx"
    book = Workbook()
    book.active.title = "Private Kunden"
    book.active.append(["Name", "Betrag"])
    book.active.append(["SecretPerson", 42.25])
    book.active["C2"] = "=B2*2"
    book.save(excel)
    word = tmp_path / "private.docx"
    doc = Document()
    doc.add_heading("Private Sektion", level=1)
    p = doc.add_paragraph("Betrag: {{")
    p.add_run("total}}")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "Kosten"
    doc.save(word)
    before = [p.read_bytes() for p in (excel, word)]
    report = inspect_documents([excel, word])
    assert report["classification"] == "CONFIDENTIAL_LOCAL_ONLY"
    assert report["production_approved"] is False
    sheet = report["files"][0]["sheets"][0]
    assert sheet["name"] == "Private Kunden"
    assert sheet["candidate_header"][0] == {"cell": "A1", "text": "Name"}
    assert sheet["formula_cells"] == 1
    assert sheet["formulas_without_cached_value"] == 1
    assert report["files"][1]["placeholders"] == ["total"]
    assert report["files"][1]["table_count"] == 1
    assert [p.read_bytes() for p in (excel, word)] == before


def test_cli_writes_only_local_report_and_never_echoes_source(tmp_path, capsys):
    source = tmp_path / "SecretPerson.xlsx"
    Workbook().save(source)
    output = tmp_path / "inventory.json"
    args = ["--input", str(source), "--output", str(output)]
    assert main(args) == 0
    first = output.read_bytes()
    assert main(args) == 20
    assert output.read_bytes() == first
    assert "SecretPerson" not in str(capsys.readouterr())


def test_blocked_file_is_reported_without_claiming_readiness(tmp_path, capsys):
    source = tmp_path / "private.xlsx"
    source.write_bytes(b"not a workbook")
    output = tmp_path / "inventory.json"
    assert main(["--input", str(source), "--output", str(output)]) == 20
    report = json.loads(output.read_text("utf-8"))
    assert report["files"][0]["status"] == "blocked"
    assert report["production_approved"] is False
    assert "private" not in str(capsys.readouterr())

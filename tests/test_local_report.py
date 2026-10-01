"""First complete local Excel-to-Word journey, including split runs and tables."""

import hashlib
import json
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from pseudokrat.ki_office import ProjectError
from pseudokrat.local_report import generate_report, main


@pytest.fixture
def report_inputs(tmp_path):
    excel = tmp_path / "original.xlsx"
    book = Workbook()
    book.active.title = "Daten"
    book.active.append(["Name", "Betrag"])
    book.active.append(["Originalname", 123.45])
    book.save(excel)
    template = tmp_path / "template.docx"
    doc = Document()
    p = doc.add_paragraph("Bericht für ")
    p.add_run("{{ na").bold = True
    p.add_run("me }}")
    doc.add_table(rows=1, cols=2).cell(0, 1).text = "{{ amount }} EUR"
    doc.sections[0].header.paragraphs[0].text = "{{ name }}"
    doc.core_properties.author = "PrivateAuthor"
    doc.save(template)
    spec = {"version": 1, "reviewed": True, "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
            "headers": {"Daten": {"A1": "Name", "B1": "Betrag"}},
            "fields": {"name": {"sheet": "Daten", "range": "A2", "operation": "cell", "format": "text"},
                       "amount": {"sheet": "Daten", "range": "B2", "operation": "cell", "format": "decimal", "decimals": 2}}}
    mapping = tmp_path / "mapping.json"
    mapping.write_text(json.dumps(spec), encoding="utf-8")
    return excel, template, mapping


def test_local_report_preserves_template_and_records_sources(tmp_path, report_inputs):
    excel, template, mapping = report_inputs
    before = [p.read_bytes() for p in report_inputs]
    output = generate_report(excel, template, mapping, tmp_path / "result")
    doc = Document(output / "bericht.docx")
    assert doc.paragraphs[0].text == "Bericht für Originalname"
    assert doc.tables[0].cell(0, 1).text == "123,45 EUR"
    assert doc.sections[0].header.paragraphs[0].text == "Originalname"
    provenance = json.loads((output / "nachweis.json").read_text("utf-8"))
    assert provenance["classification"] == "CONFIDENTIAL_LOCAL_ONLY"
    assert provenance["production_approved"] is False
    assert provenance["facts"]["amount"]["exact_value"] == "123.45"
    assert provenance["output_sha256"] == hashlib.sha256((output / "bericht.docx").read_bytes()).hexdigest()
    with ZipFile(output / "bericht.docx") as archive:
        assert not any(n.startswith("docProps/") for n in archive.namelist())
    assert [p.read_bytes() for p in report_inputs] == before
    with pytest.raises(ProjectError):
        generate_report(excel, template, mapping, output)


def test_missing_mapping_field_leaves_no_output(tmp_path, report_inputs):
    excel, template, mapping = report_inputs
    spec = json.loads(mapping.read_text("utf-8"))
    del spec["fields"]["amount"]
    mapping.write_text(json.dumps(spec), encoding="utf-8")
    with pytest.raises(ProjectError):
        generate_report(excel, template, mapping, tmp_path / "result")
    assert not (tmp_path / "result").exists()


def test_cli_does_not_echo_original_values(tmp_path, report_inputs, capsys):
    excel, template, mapping = report_inputs
    assert main(["--excel", str(excel), "--template", str(template), "--mapping", str(mapping),
                 "--output", str(tmp_path / "result")]) == 0
    assert "Originalname" not in str(capsys.readouterr())


def test_sixty_five_section_report(tmp_path, report_inputs):
    excel, template, mapping = report_inputs
    doc = Document()
    for i in range(65):
        if i:
            doc.add_page_break()
        doc.add_heading(f"Abschnitt {i + 1}", level=1)
        doc.add_paragraph("Bericht für {{ name }}")
        doc.add_table(rows=1, cols=2).cell(0, 1).text = "{{ amount }} EUR"
    doc.save(template)
    spec = json.loads(mapping.read_text("utf-8"))
    spec["template_sha256"] = hashlib.sha256(template.read_bytes()).hexdigest()
    mapping.write_text(json.dumps(spec), encoding="utf-8")
    output = generate_report(excel, template, mapping, tmp_path / "result")
    result = Document(output / "bericht.docx")
    assert len(result.tables) == 65
    assert all(t.cell(0, 1).text == "123,45 EUR" for t in result.tables)
    assert sum(p.text == "Bericht für Originalname" for p in result.paragraphs) == 65

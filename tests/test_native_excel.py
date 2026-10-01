"""Real Excel tests are explicit opt-in; input gates run without Excel."""

import os

import pytest
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule

from pseudokrat.ki_office import ProjectError
from pseudokrat.native_excel import recalculate, validate_source


def test_unsafe_formula_never_reaches_excel(tmp_path):
    source = tmp_path / "input.xlsx"
    book = Workbook()
    book.active["A1"] = '=WEBSERVICE("https://example.invalid")'
    book.save(source)
    with pytest.raises(ProjectError):
        validate_source(source)


def test_unsafe_conditional_formula_never_reaches_excel(tmp_path):
    source = tmp_path / "conditional.xlsx"
    book = Workbook()
    book.active["A1"] = 1
    book.active.conditional_formatting.add("A1", FormulaRule(formula=['WEBSERVICE("https://example.invalid")']))
    book.save(source)
    with pytest.raises(ProjectError):
        validate_source(source)


@pytest.mark.parametrize("setting,value", [("iterate", True), ("fullPrecision", False)])
def test_calculation_semantics_are_not_silently_changed(tmp_path, setting, value):
    source = tmp_path / "special.xlsx"
    book = Workbook()
    book.active["A1"] = "=1+2"
    setattr(book.calculation, setting, value)
    book.save(source)
    with pytest.raises(ProjectError):
        validate_source(source)


@pytest.mark.skipif(os.environ.get("PSEUDOKRAT_TEST_EXCEL") != "1", reason="Requires explicit local Excel opt-in")
def test_real_excel_rebuild_returns_fresh_values_without_modifying_original(tmp_path):
    source = tmp_path / "input.xlsx"
    book = Workbook()
    book.active.title = "Daten"
    book.active["A1"] = 12.35
    book.active["A2"] = -2.15
    book.active["B1"] = "=SUM(A1:A2)"
    book.active["B2"] = "=ROUND(-0.125,2)"
    book.active["B3"] = '=IF(A1>0,"positiv","negativ")'
    hidden = book.create_sheet("Versteckt")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "=Daten!A1*2"
    book.save(source)
    original = source.read_bytes()
    result = recalculate(source)
    assert result["engine"] == "Microsoft Excel"
    assert result["cells"]["Daten"]["B1"]["value"] == pytest.approx(10.2)
    assert result["cells"]["Daten"]["B2"]["value"] == pytest.approx(-0.13)
    assert result["cells"]["Daten"]["B3"]["value"] == "positiv"
    assert result["cells"]["Versteckt"]["A1"]["value"] == pytest.approx(24.7)
    assert source.read_bytes() == original


@pytest.mark.skipif(os.environ.get("PSEUDOKRAT_TEST_EXCEL") != "1", reason="Requires explicit local Excel opt-in")
def test_real_excel_error_cell_stops_result(tmp_path):
    source = tmp_path / "error.xlsx"
    book = Workbook()
    book.active["A1"] = "=1/0"
    book.save(source)
    with pytest.raises(ProjectError):
        recalculate(source)


@pytest.mark.skipif(os.environ.get("PSEUDOKRAT_TEST_EXCEL") != "1", reason="Requires explicit local Excel opt-in")
@pytest.mark.parametrize("formula", ["=A1+1", "=B1+1"])
def test_real_excel_circular_reference_stops_result(tmp_path, formula):
    source = tmp_path / "circular.xlsx"
    book = Workbook()
    book.active["A1"] = formula
    book.active["B1"] = "=A1+1"
    book.save(source)
    with pytest.raises(ProjectError):
        recalculate(source)

"""Opt-in real Linux XLSX calculation tests; use synthetic sources only."""

import json
import os
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from pseudokrat import native_libreoffice
from pseudokrat.ki_office import ProjectError
from pseudokrat.native_libreoffice import recalculate

pytestmark = pytest.mark.skipif(sys.platform != "linux" or os.environ.get("PSEUDOKRAT_TEST_LIBREOFFICE") != "1",
                               reason="Requires explicit Linux LibreOffice opt-in")


def test_real_xlsx_calculation_preserves_original(tmp_path):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "Daten"
    book.active["A1"] = 12.35
    book.active["A2"] = -2.15
    book.active["B1"] = "=SUM(A1:A2)"
    book.active["B2"] = "=ROUND(-0.125,2)"
    book.active["B3"] = '=IF(A1>0,"positiv","negativ")'
    book.active["B4"] = "=IF(A1>0,7,-7)"
    sheet = book.create_sheet("Versteckt")
    sheet.sheet_state = "veryHidden"
    sheet["A1"] = "=Daten!A1*2"
    book.save(path)
    original = path.read_bytes()
    result = recalculate(path)
    assert result["engine"] == "LibreOffice Calc"
    assert result["cells"]["Daten"]["B1"]["value"] == pytest.approx(10.2)
    assert result["cells"]["Daten"]["B2"]["value"] == pytest.approx(-0.13)
    assert result["cells"]["Daten"]["B3"]["value"] == "positiv"
    assert result["cells"]["Daten"]["B4"]["value"] == 7
    assert result["cells"]["Versteckt"]["A1"]["value"] == pytest.approx(24.7)
    assert path.read_bytes() == original


@pytest.mark.parametrize("formula", ["=1/0", "=A1+1", '=WEBSERVICE("https://example.invalid")'])
def test_error_and_unsafe_formulas_are_blocked(tmp_path, formula):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active["A1"] = formula
    book.save(path)
    with pytest.raises(ProjectError):
        recalculate(path)


@pytest.mark.parametrize("formula", ["=TRUE()", "=1=1", "=IF(1,TRUE(),FALSE())", "=--TRUE()"])
def test_boolean_formula_is_not_a_monetary_value(tmp_path, formula):
    path = tmp_path / "boolean.xlsx"
    book = Workbook()
    book.active["A1"] = formula
    book.active["A1"].number_format = "0.00"
    book.save(path)
    with pytest.raises(ProjectError):
        recalculate(path)


def test_boolean_source_reference_is_blocked(tmp_path):
    path = tmp_path / "boolean_source.xlsx"
    book = Workbook()
    book.active["A1"] = True
    book.active["A1"].number_format = "0.00"
    book.active["B1"] = "=A1"
    book.save(path)
    with pytest.raises(ProjectError):
        recalculate(path)


def test_masked_boolean_dependency_blocks_entire_workbook(tmp_path):
    path = tmp_path / "masked.xlsx"
    book = Workbook()
    book.active["A1"] = "=TRUE()"
    book.active["A1"].number_format = "0.00"
    book.active["B1"] = "=A1"
    book.active["C1"] = "=IFERROR(B1,42)"
    book.save(path)
    with pytest.raises(ProjectError):
        recalculate(path)


@pytest.mark.parametrize("formula", ["=RAND()", "=IF(RAND()>0.5,TRUE(),1)", "=TODAY()"])
def test_volatile_formulas_are_not_evaluated_twice(tmp_path, formula):
    path = tmp_path / "volatile.xlsx"
    book = Workbook()
    book.active["A1"] = formula
    book.save(path)
    with pytest.raises(ProjectError):
        recalculate(path)


def test_timeout_cleans_working_copy(tmp_path, monkeypatch):
    path = tmp_path / "timeout.xlsx"
    book = Workbook()
    book.active["A1"] = "=1+2"
    book.save(path)
    original = path.read_bytes()
    marker = tmp_path / "ownership.json"
    worker = Path(native_libreoffice.__file__).with_name("libreoffice_worker.py").read_text()
    worker = worker.replace("        document.calculateAll()", f"        Path({str(marker)!r}).write_text(json.dumps({{'folder': str(folder), 'group': __import__('os').getpgrp()}}))\n        time.sleep(60)\n        document.calculateAll()")
    (tmp_path / "libreoffice_worker.py").write_text(worker)
    monkeypatch.setattr(native_libreoffice, "__file__", str(tmp_path / "native_libreoffice.py"))
    with pytest.raises(ProjectError):
        recalculate(path, timeout=15)
    assert marker.exists()
    ownership = json.loads(marker.read_text())
    assert not Path(ownership["folder"]).exists()
    for entry in Path("/proc").iterdir():
        if entry.name.isdigit():
            try:
                status = (entry / "stat").read_text().rsplit(")", 1)[1].split()
            except (FileNotFoundError, ProcessLookupError):
                continue
            if int(status[2]) == ownership["group"]:
                assert status[0] == "Z", "An owned process remains running"
    assert path.read_bytes() == original

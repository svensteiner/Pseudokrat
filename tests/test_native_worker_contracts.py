"""Untrusted worker results must fail closed even without installed Office."""

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from pseudokrat import native_libreoffice, report_preview
from pseudokrat.ki_office import ProjectError


@pytest.fixture
def sources(tmp_path):
    excel = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "Daten"
    book.active["A1"] = 12
    book.active["B1"] = "=A1*2"
    book.save(excel)
    word = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Synthetischer Bericht")
    doc.save(word)
    return excel, word


def worker(monkeypatch, module, emit, *, code=0):
    """Simulate only the OS boundary; real parsers and validators still run."""
    folders, killed = [], []

    class Process:
        pid = 987654

        def __init__(self, command, **options):
            assert options["start_new_session"] is True
            folder = Path(command[-1])
            folders.append(folder)
            emit(folder)

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def wait(self, timeout):
            return code

    monkeypatch.setattr(module.platform, "system", lambda: "Linux")
    monkeypatch.setattr(module.subprocess, "Popen", Process)
    monkeypatch.setattr(module.os, "killpg", lambda pid, sig: killed.append(pid), raising=False)
    monkeypatch.setattr(module.signal, "SIGKILL", 9, raising=False)
    return folders, killed


def calc_result():
    return {
        "engine": "LibreOffice Calc",
        "version": "test",
        "cells": {"Daten": {"B1": {"value": 24}}},
    }


@pytest.mark.parametrize(
    "fault",
    [
        "engine",
        "version",
        "sheet",
        "cell",
        "shape",
        "bool",
        "empty",
        "nan",
        "json",
        "source",
        "exit",
    ],
)
def test_calc_rejects_invalid_worker_evidence(tmp_path, monkeypatch, sources, fault):
    excel, _ = sources
    original = excel.read_bytes()

    def emit(folder):
        result = calc_result()
        if fault == "engine":
            result["engine"] = "unknown"
        elif fault == "version":
            result["version"] = None
        elif fault == "sheet":
            result["cells"]["extra"] = {}
        elif fault == "cell":
            result["cells"]["Daten"]["C1"] = {"value": 24}
        elif fault == "shape":
            result["cells"]["Daten"]["B1"] = 24
        elif fault in {"bool", "empty", "nan"}:
            result["cells"]["Daten"]["B1"]["value"] = {
                "bool": True,
                "empty": "",
                "nan": float("nan"),
            }[fault]
        elif fault == "source":
            excel.write_bytes(original + b"changed")
        (folder / "result.json").write_text(
            "{" if fault == "json" else json.dumps(result), encoding="utf-8"
        )

    folders, killed = worker(
        monkeypatch, native_libreoffice, emit, code=20 if fault == "exit" else 0
    )
    with pytest.raises(ProjectError):
        native_libreoffice.recalculate(excel)
    assert killed == [987654]
    assert all(not path.exists() for path in folders)
    if fault != "source":
        assert excel.read_bytes() == original


def test_calc_binds_valid_worker_result_to_original_snapshot(monkeypatch, sources):
    excel, _ = sources
    original = excel.read_bytes()

    def emit(folder):
        assert json.loads((folder / "request.json").read_text()) == {"Daten": ["B1"]}
        assert (folder / "input.xlsx").exists()
        (folder / "result.json").write_text(json.dumps(calc_result()), encoding="utf-8")

    folders, killed = worker(monkeypatch, native_libreoffice, emit)
    result = native_libreoffice.recalculate(excel)
    assert result["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert result["cells"]["Daten"]["B1"]["value"] == 24
    assert excel.read_bytes() == original
    assert killed == [987654]
    assert all(not path.exists() for path in folders)


@pytest.mark.parametrize(
    "fault", [None, "pdf", "empty", "encrypted", "engine", "json", "source", "exit"]
)
def test_preview_validates_output_before_publishing(tmp_path, monkeypatch, sources, fault):
    _, word = sources
    original = word.read_bytes()
    destination = tmp_path / "preview"

    def emit(folder):
        pdf = PdfWriter()
        if fault != "empty":
            pdf.add_blank_page(width=100, height=100)
        if fault == "encrypted":
            pdf.encrypt("synthetic-password")
        pdf.write(folder / "preview.pdf")
        if fault == "pdf":
            (folder / "preview.pdf").write_bytes(b"not a PDF")
        engine = {
            "engine": "unknown" if fault == "engine" else "LibreOffice Writer",
            "version": "test",
        }
        (folder / "engine.json").write_text(
            "{" if fault == "json" else json.dumps(engine), encoding="utf-8"
        )
        if fault == "source":
            word.write_bytes(original + b"changed")

    folders, killed = worker(monkeypatch, report_preview, emit, code=20 if fault == "exit" else 0)
    if fault:
        with pytest.raises(ProjectError):
            report_preview.render_preview(word, destination)
        assert not destination.exists()
    else:
        report_preview.render_preview(word, destination)
        evidence = json.loads((destination / "preview.json").read_text())
        assert evidence["source_sha256"] == hashlib.sha256(original).hexdigest()
        assert (
            evidence["pdf_sha256"]
            == hashlib.sha256((destination / "preview.pdf").read_bytes()).hexdigest()
        )
        assert evidence["pages"] == 1
        assert evidence["production_approved"] is False
        assert evidence["visual_reviewed"] is False
    assert killed == [987654]
    assert all(not path.exists() for path in folders)
    if fault != "source":
        assert word.read_bytes() == original

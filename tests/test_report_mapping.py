"""The source-to-report contract must fail on drift, missing data and stale formulas."""

import hashlib

import pytest
from docx import Document
from openpyxl import Workbook

from pseudokrat.ki_office import ProjectError
from pseudokrat.report_mapping import resolve_mapping


@pytest.fixture
def sources(tmp_path):
    excel = tmp_path / "input.xlsx"
    book = Workbook()
    book.active.title = "Daten"
    book.active.append(["Name", "Betrag"])
    book.active.append(["Beispiel", 12.35])
    book.active.append(["Beispiel", -2.15])
    book.save(excel)
    template = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("{{ name }}: {{ total }} EUR")
    doc.save(template)
    spec = {
        "version": 1,
        "reviewed": True,
        "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
        "headers": {"Daten": {"A1": "Name", "B1": "Betrag"}},
        "fields": {
            "name": {"sheet": "Daten", "range": "A2", "operation": "cell", "format": "text"},
            "total": {
                "sheet": "Daten",
                "range": "B2:B3",
                "operation": "sum",
                "format": "decimal",
                "decimals": 2,
            },
        },
    }
    return excel, template, spec


def test_exact_decimal_sum_with_provenance(sources):
    excel, template, spec = sources
    result = resolve_mapping(excel, template, spec)
    assert result["name"]["text"] == "Beispiel"
    assert result["total"]["text"] == "10,20"
    assert result["total"]["exact_value"] == "10.20"
    assert result["total"]["source"] == {"sheet": "Daten", "range": "B2:B3", "operation": "sum"}


@pytest.mark.parametrize(
    "change", ["review", "template", "header", "missing", "formula", "unknown", "integer"]
)
def test_untrusted_or_ambiguous_mapping_blocks(sources, change):
    excel, template, spec = sources
    if change == "review":
        spec["reviewed"] = False
    elif change == "template":
        spec["template_sha256"] = "0" * 64
    elif change == "unknown":
        spec["fields"]["total"]["python"] = "arbitrary code"
    elif change == "integer":
        spec["fields"]["total"]["format"] = "integer"
        spec["fields"]["total"].pop("decimals")
    else:
        from openpyxl import load_workbook

        book = load_workbook(excel)
        if change == "header":
            book.active["B1"] = "Andere Einheit"
        elif change == "missing":
            book.active["B3"] = None
        else:
            book.active["B3"] = "=1+1"
        book.save(excel)
    with pytest.raises(ProjectError):
        resolve_mapping(excel, template, spec)


@pytest.mark.parametrize(
    "value,number_format,expected",
    [
        (0, "General", "0,00"),
        (-0.125, "General", "-0,13"),
        (45292, "yyyy-mm-dd", None),
        (True, "General", None),
    ],
)
def test_rounding_zero_dates_and_boolean_types(sources, value, number_format, expected):
    from openpyxl import load_workbook

    excel, template, spec = sources
    book = load_workbook(excel)
    book.active["B2"] = value
    book.active["B2"].number_format = number_format
    book.save(excel)
    spec["fields"]["total"]["range"] = "B2"
    if expected is None:
        with pytest.raises(ProjectError):
            resolve_mapping(excel, template, spec)
    else:
        assert resolve_mapping(excel, template, spec)["total"]["text"] == expected


def test_unknown_locale_specific_number_format_blocks(sources):
    from pseudokrat.ki_office import S, read_office, write_office

    excel, template, spec = sources
    roots = read_office(excel)
    xf = roots["xl/styles.xml"].find(f"{{{S}}}cellXfs")[0]
    xf.set("numFmtId", "27")
    excel.write_bytes(write_office(roots))
    with pytest.raises(ProjectError):
        resolve_mapping(excel, template, spec)

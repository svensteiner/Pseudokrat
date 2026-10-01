"""Opt-in native Writer rendering with actual page-count field verification."""

import hashlib
import json
import os
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

from pseudokrat import report_preview
from pseudokrat.ki_office import ProjectError
from pseudokrat.report_preview import render_preview

pytestmark = pytest.mark.skipif(
    sys.platform != "linux" or os.environ.get("PSEUDOKRAT_TEST_LIBREOFFICE") != "1",
    reason="Requires explicit Linux LibreOffice opt-in",
)


def field(paragraph, code, locked=False):
    node = OxmlElement("w:fldSimple")
    node.set(qn("w:instr"), code)
    if locked:
        node.set(qn("w:fldLock"), "true")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "999"
    run.append(text)
    node.append(run)
    paragraph._p.append(node)


@pytest.mark.parametrize("pages", [3, 65])
def test_real_writer_updates_preview_pagination_without_changing_docx(tmp_path, pages):
    source = tmp_path / "report.docx"
    doc = Document()
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run("Seite ")
    field(footer, "PAGE", locked=pages == 65)
    footer.add_run(" von ")
    field(footer, "NUMPAGES", locked=pages == 65)
    for index in range(pages):
        if index:
            doc.add_page_break()
        doc.add_heading(f"Abschnitt {index + 1}", level=1)
        doc.add_paragraph("Synthetischer Bericht. Saldo: 24,69 EUR.")
    doc.save(source)
    original = source.read_bytes()
    output = render_preview(source, tmp_path / "preview")
    pdf = PdfReader(output / "preview.pdf")
    assert len(pdf.pages) == pages
    first, last = (" ".join(page.extract_text().split()) for page in (pdf.pages[0], pdf.pages[-1]))
    assert f"Seite 1 von {pages}" in first
    assert f"Seite {pages} von {pages}" in last
    assert "999" not in first + last
    evidence = json.loads((output / "preview.json").read_text("utf-8"))
    assert evidence["pages"] == pages
    assert evidence["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert evidence["source_docx_updated"] is False
    assert evidence["production_approved"] is False
    assert source.read_bytes() == original
    with pytest.raises(ProjectError):
        render_preview(source, output)


def test_unfilled_report_is_blocked_before_native_open(tmp_path, monkeypatch):
    source = tmp_path / "unfilled.docx"
    doc = Document()
    doc.add_paragraph("{{ missing }}")
    doc.save(source)

    def no_start(*args, **kwargs):
        raise AssertionError("Native worker must not start")

    monkeypatch.setattr("pseudokrat.report_preview.subprocess.Popen", no_start)
    with pytest.raises(ProjectError):
        render_preview(source, tmp_path / "preview")
    assert not (tmp_path / "preview").exists()


def test_writer_timeout_leaves_no_preview_or_working_copy(tmp_path, monkeypatch):
    source = tmp_path / "timeout.docx"
    doc = Document()
    doc.add_paragraph("Synthetischer Test")
    doc.save(source)
    original = source.read_bytes()
    marker = tmp_path / "ownership.json"
    worker = Path(report_preview.__file__).with_name("writer_preview_worker.py").read_text()
    worker = worker.replace(
        "        document.reformat()",
        f"        Path({str(marker)!r}).write_text(json.dumps({{'folder': str(folder), 'group': __import__('os').getpgrp()}}))\n        time.sleep(60)\n        document.reformat()",
        1,
    )
    (tmp_path / "writer_preview_worker.py").write_text(worker)
    monkeypatch.setattr(report_preview, "__file__", str(tmp_path / "report_preview.py"))
    with pytest.raises(ProjectError):
        render_preview(source, tmp_path / "preview", timeout=15)
    ownership = json.loads(marker.read_text())
    assert not Path(ownership["folder"]).exists()
    assert not (tmp_path / "preview").exists()
    assert source.read_bytes() == original
    for entry in Path("/proc").iterdir():
        if entry.name.isdigit():
            try:
                status = (entry / "stat").read_text().rsplit(")", 1)[1].split()
            except (FileNotFoundError, ProcessLookupError):
                continue
            if int(status[2]) == ownership["group"]:
                assert status[0] == "Z"

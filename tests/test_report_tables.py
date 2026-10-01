"""Dynamic rows must preserve values, formatting and Excel row provenance."""

import hashlib
import json

import pytest
from docx import Document
from openpyxl import Workbook

from pseudokrat.ki_office import ProjectError
from pseudokrat.local_report import generate_report


@pytest.fixture
def table_inputs(tmp_path):
    excel = tmp_path / "input.xlsx"
    book = Workbook()
    book.active.title = "Daten"
    book.active.append(["Name", "Betrag"])
    book.active.append(["Alpha", 0])
    book.active.append(["Beta", -12.35])
    book.active.append(["Gamma", 123.45])
    book.save(excel)
    template = tmp_path / "template.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Betrag EUR"
    table.cell(1, 0).paragraphs[0].add_run("{{ rows.name }}").bold = True
    table.cell(1, 1).text = "{{ rows.amount }}"
    table.cell(2, 0).text = "Ende"
    doc.save(template)
    spec = {
        "version": 2,
        "reviewed": True,
        "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
        "headers": {"Daten": {"A1": "Name", "B1": "Betrag"}},
        "fields": {},
        "tables": {
            "rows": {
                "sheet": "Daten",
                "first_row": 2,
                "last_row": 4,
                "columns": {
                    "name": {"column": "A", "format": "text"},
                    "amount": {"column": "B", "format": "decimal", "decimals": 2},
                },
            }
        },
    }
    mapping = tmp_path / "mapping.json"
    mapping.write_text(json.dumps(spec), encoding="utf-8")
    return excel, template, mapping


def test_rows_expand_between_existing_header_and_footer(tmp_path, table_inputs):
    output = generate_report(*table_inputs, tmp_path / "output")
    table = Document(output / "bericht.docx").tables[0]
    assert len(table.rows) == 5
    assert [table.cell(i, 0).text for i in range(5)] == ["Name", "Alpha", "Beta", "Gamma", "Ende"]
    assert [table.cell(i, 1).text for i in range(1, 4)] == ["0,00", "-12,35", "123,45"]
    assert all(table.cell(i, 0).paragraphs[0].runs[0].bold for i in range(1, 4))
    evidence = json.loads((output / "nachweis.json").read_text("utf-8"))
    assert evidence["tables"]["rows"][1]["amount"]["source"]["range"] == "B3"


def test_one_hundred_twenty_rows_preserve_order_and_provenance(tmp_path, table_inputs):
    from openpyxl import load_workbook

    excel, template, mapping = table_inputs
    book = load_workbook(excel)
    for i in range(3, 120):
        book.active.append([f"Item {i}", i + 0.25])
    book.save(excel)
    spec = json.loads(mapping.read_text("utf-8"))
    spec["tables"]["rows"]["last_row"] = 121
    mapping.write_text(json.dumps(spec), encoding="utf-8")
    output = generate_report(excel, template, mapping, tmp_path / "result")
    table = Document(output / "bericht.docx").tables[0]
    assert len(table.rows) == 122
    assert table.cell(120, 0).text == "Item 119"
    assert table.cell(120, 1).text == "119,25"
    evidence = json.loads((output / "nachweis.json").read_text("utf-8"))
    assert len(evidence["tables"]["rows"]) == 120
    assert evidence["tables"]["rows"][-1]["amount"]["source"]["range"] == "B121"


def test_table_internal_keys_do_not_collide():
    from pseudokrat.report_tables import compile_tables

    spec = {
        "version": 2,
        "reviewed": True,
        "template_sha256": "0" * 64,
        "headers": {},
        "fields": {},
        "tables": {
            "a_2": {
                "sheet": "Daten",
                "first_row": 3,
                "last_row": 3,
                "columns": {"b": {"column": "A", "format": "text"}},
            },
            "a": {
                "sheet": "Daten",
                "first_row": 2,
                "last_row": 2,
                "columns": {"_3_b": {"column": "B", "format": "text"}},
            },
        },
    }
    compiled, groups = compile_tables(spec)
    assert len(compiled["fields"]) == 2
    assert groups["a_2"][0]["b"] != groups["a"][0]["_3_b"]


@pytest.mark.parametrize("case", ["missing", "outside", "merge", "bookmark", "malformed", "prefix"])
def test_ambiguous_rows_and_incomplete_sources_block(tmp_path, table_inputs, case):
    excel, template, mapping = table_inputs
    if case == "missing":
        from openpyxl import load_workbook

        book = load_workbook(excel)
        book.active["B3"] = None
        book.save(excel)
    else:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document(template)
        if case == "outside":
            doc.add_paragraph("{{ rows.name }}")
        elif case == "merge":
            doc.tables[0].cell(1, 0).merge(doc.tables[0].cell(2, 0))
        elif case == "bookmark":
            element = OxmlElement("w:bookmarkStart")
            element.set(qn("w:id"), "1")
            element.set(qn("w:name"), "private")
            doc.tables[0].cell(1, 0).paragraphs[0]._p.append(element)
        elif case == "malformed":
            doc.tables[0].cell(1, 1).text = "{{ rows.amount }} {{broken"
        else:
            doc.tables[0].cell(1, 0).text = "{{ rows.unknown }}"
        doc.save(template)
        spec = json.loads(mapping.read_text("utf-8"))
        spec["template_sha256"] = hashlib.sha256(template.read_bytes()).hexdigest()
        mapping.write_text(json.dumps(spec), encoding="utf-8")
    with pytest.raises(ProjectError):
        generate_report(excel, template, mapping, tmp_path / "output")
    assert not (tmp_path / "output").exists()

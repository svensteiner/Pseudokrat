"""Safe pagination fields remain local and cannot hide executable instructions."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pseudokrat.ki_office import ProjectError, read_local_template, read_office
from pseudokrat.word_fields import inspect_fields, text_segments


def complex_field(paragraph, instructions, cached="1", close=True):
    run = paragraph.add_run()._r
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    run.append(start)
    for instruction in instructions:
        node = OxmlElement("w:instrText")
        node.text = instruction
        paragraph.add_run()._r.append(node)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separator)
    paragraph.add_run(cached)
    if close:
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        paragraph.add_run()._r.append(end)


@pytest.mark.parametrize(
    "code", ["PAGE", "NUMPAGES", "SECTION", "SECTIONPAGES", r"PAGE \* MERGEFORMAT"]
)
def test_local_pagination_is_preserved_but_export_stays_blocked(tmp_path, code):
    path = tmp_path / "template.docx"
    doc = Document()
    p = doc.sections[0].footer.paragraphs[0]
    p.add_run("{{ name }} Seite ")
    complex_field(p, [code[:2], code[2:]])
    doc.save(path)
    roots, _ = read_local_template(path)
    codes, protected = inspect_fields(roots)
    assert codes == [code.split()[0]]
    assert [node.text for node in protected] == ["1"]
    with pytest.raises(ProjectError):
        read_office(path)


@pytest.mark.parametrize(
    "instruction",
    [
        "DDEAUTO cmd",
        'INCLUDETEXT "https://example.invalid"',
        'INCLUDEPICTURE "file:///secret"',
        "PAGE DDE",
        r"PAGE \x unknown",
        "TOC",
    ],
)
def test_other_fields_are_blocked_even_with_split_instruction(tmp_path, instruction):
    path = tmp_path / "bad.docx"
    doc = Document()
    complex_field(doc.add_paragraph(), [instruction[:3], instruction[3:]])
    doc.save(path)
    with pytest.raises(ProjectError):
        read_local_template(path)


def test_unclosed_field_is_blocked(tmp_path):
    path = tmp_path / "bad.docx"
    doc = Document()
    complex_field(doc.add_paragraph(), ["PAGE"], close=False)
    doc.save(path)
    with pytest.raises(ProjectError):
        read_local_template(path)


def test_placeholder_cannot_cross_field_boundary():
    doc = Document()
    p = doc.add_paragraph("{{ name")
    complex_field(p, ["PAGE"])
    p.add_run(" }}")
    _, protected = inspect_fields({"part": doc._element})
    assert ["".join(n.text for n in segment) for segment in text_segments(p._p, protected)] == [
        "{{ name",
        " }}",
    ]


def test_placeholder_inside_field_result_is_rejected():
    doc = Document()
    complex_field(doc.add_paragraph(), ["PAGE"], "{{ name }}")
    with pytest.raises(ProjectError):
        inspect_fields({"part": doc._element})


def test_nested_complex_fields_are_rejected():
    doc = Document()
    p = doc.add_paragraph()
    complex_field(p, ["PAGE"], close=False)
    complex_field(p, ["NUMPAGES"])
    with pytest.raises(ProjectError):
        inspect_fields({"part": doc._element})
